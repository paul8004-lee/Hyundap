"""ingest/loaders/docx_loader.py — DOCX 문단·표를 공통 섹션으로 변환.

담당: 김혜리
기준 명세: AI_문서검색_통합_기능명세서

함수:
    FUN-209 load_docx — DOCX 문단·표를 공통 문서로 읽는다. 그림은 PaddleOCR(있으면)로, Office 차트는 차트 XML에서 읽는다."""


import re
import zlib
from zipfile import BadZipFile

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml.etree import XMLSyntaxError

from . import charts, figures
from .base import build_location_label, table_to_text
from hyundap.common.schemas import CommonDocument, Section
from ..models import CORRUPTED_FILE, IngestError

# 글자를 읽지 않는 그림·도형·차트·텍스트 상자(w:drawing, 구형 w:pict). mc:Fallback 안의 중복 표현은 뺀다.
UNREAD_OBJECTS_XPATH = ".//w:drawing | .//w:pict[not(ancestor::*[local-name()='Fallback'])]"
EMU_PER_PT = 12700

# 제목 스타일(Heading)이 문서에 하나도 없을 때만 쓰는 보조 제목 규칙. 한국 사무 문서는 제목을 Normal 스타일에 번호만 붙여 쓰는 경우가 많다.
# 큰 번호 제목만 새 섹션으로 본다: '제1장', '1.', 'Ⅰ.', 'I.'. '1)'·'가.' 같은 작은 번호는 본문으로 둔다(작은 섹션이 많아지는 것을 막는다).
NUMBERED_HEADING = re.compile(r"^(?:제\s*\d+\s*[장절편]|\d{1,2}\.(?!\d)|[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]+\.|[IVX]+\.)\s*\S")
HEADING_MAX_CHARS = 40   # 이보다 긴 문단은 제목이 아니라 번호가 붙은 문장으로 본다
SENTENCE_ENDINGS = (".", "!", "?", "다")   # 이 끝말이면 제목이 아니라 문장이다('개요'처럼 '요'로 끝나는 제목은 통과)


def style_name_of(block, style_names):
    """문단의 스타일 이름. 같은 스타일 번호는 한 번만 읽어 style_names에 저장한다(이름을 읽는 것이 느리다)."""
    style_id = block._p.style
    if style_id not in style_names:
        style_names[style_id] = block.style.name
    return style_names[style_id]


def has_heading_style(docx_file, style_names):
    """본문에 제목 스타일(Heading) 문단이 하나라도 있으면 True."""
    for block in iter_blocks(docx_file.element.body, docx_file._body):
        if isinstance(block, Paragraph) and style_name_of(block, style_names).startswith("Heading"):
            return True
    return False


def looks_like_heading(text, style_name, paragraph):
    """제목 스타일이 없는 문서에서 이 문단을 제목으로 볼지 정한다.

    아웃라인 수준(w:outlineLvl)이 있으면 제목이다. 아니면 짧고(HEADING_MAX_CHARS) 문장으로 끝나지 않고
    큰 번호(NUMBERED_HEADING)로 시작하는 문단만 제목이다. 목록 문단은 제외한다."""
    if text == "":
        return False
    properties = paragraph._p.pPr
    if properties is not None and properties.find(qn("w:outlineLvl")) is not None:
        return True
    if style_name.startswith("List"):
        return False
    return len(text) <= HEADING_MAX_CHARS and not text.endswith(SENTENCE_ENDINGS) and NUMBERED_HEADING.match(text) is not None


def add_docx_section(doc, title, lines):
    """제목과 모아 둔 글로 구간 1개를 만들어 문서에 더한다."""
    if title is None and len(lines) == 0:
        return
    section_no = len(doc.sections) + 1
    texts = []
    if title is not None and title != "":
        texts.append(title)
    texts.extend(lines)
    doc.sections.append(
        Section(
            section_no=section_no,
            location_label=build_location_label("section", section_no, title),
            text="\n\n".join(texts),
        )
    )


def emu_to_pt(value):
    """wp:extent의 cx·cy(EMU)를 pt로 바꾼다. 없거나 정수가 아니면 None(크기를 모르는 그림으로 다룬다)."""
    try:
        return int(value) / EMU_PER_PT
    except (TypeError, ValueError):
        return None


def picture_groups(element):
    """요소 안 개체(w:drawing, 구형 w:pict)마다 그 안 그림의 (관계 id, 표시 너비 pt, 표시 높이 pt) 목록을 돌려준다.

    구형 w:pict 그림은 크기를 몰라 None이다. 그림이 없는 개체(도형·차트)는 빈 목록이다."""
    groups = []
    for drawing in element.xpath(".//w:drawing[not(ancestor::*[local-name()='Fallback'])]"):
        extent = drawing.xpath(".//wp:extent")
        width = height = None
        if extent:
            width, height = emu_to_pt(extent[0].get("cx")), emu_to_pt(extent[0].get("cy"))
            if width is None or height is None:
                width = height = None
        groups.append([(rid, width, height) for rid in drawing.xpath(".//a:blip/@r:embed")])
    for pict in element.xpath(".//w:pict[not(ancestor::*[local-name()='Fallback'])]"):
        groups.append([(rid, None, None) for rid in pict.xpath(".//*[local-name()='imagedata']/@*[local-name()='id']")])
    return groups


def picture_refs(element):
    """요소 안 그림의 (관계 id, 표시 너비 pt, 표시 높이 pt) 목록. 구형 w:pict 그림은 크기를 몰라 None이다."""
    return [ref for group in picture_groups(element) for ref in group]


def read_charts(docx_file, chart_reader, element):
    """요소 안 Office 차트들에서 읽은 표 조각 목록을 돌려준다. 차트 XML은 문서 관계(r:id)로 찾는다."""
    texts = []
    for rid in element.xpath(".//w:drawing[not(ancestor::*[local-name()='Fallback'])]//c:chart/@r:id"):
        try:
            part = docx_file.part.related_parts.get(rid)   # 외부 연결 차트 등은 관계가 없다
            if part is not None:
                texts.extend(chart_reader.read(part.blob))
        except Exception:   # 차트 하나를 꺼내지 못해도 업로드 전체가 실패하지 않게 건너뛴 차트로 센다
            chart_reader.unreadable += 1
    return texts


def read_pictures(docx_file, reader, element, lines):
    """요소 안 그림들에서 OCR로 읽은 글자 목록을 돌려준다. 같은 구간에 이미 모은 글(lines)과 겹치는 줄은 뺀다."""
    texts = []   # OCR이 꺼져 있어도 그림을 훑어 읽지 못한 그림 수를 센다(read_picture가 센다)
    for group in picture_groups(element):
        read_any = False
        for rid, width, height in group:
            try:
                part = docx_file.part.related_parts.get(rid)   # 외부 연결 그림 등은 관계가 없다
                if part is None:
                    continue
                text = reader.read_picture(part.blob, width, height, page_text=lambda: "\n".join(lines + texts))
            except Exception:   # 그림 하나를 꺼내지 못해도 업로드 전체가 실패하지 않게 건너뛴 그림으로 센다
                reader.unreadable += 1
                continue
            if text != "":
                texts.append(text)
                read_any = True
        if read_any:
            reader.objects_read += 1   # 개체 하나에 그림이 여러 장 있어도 읽은 개체는 1개다
    return texts


def iter_blocks(element, parent):
    """본문 요소를 순서대로 Paragraph·Table로 돌려준다. 콘텐츠 컨트롤(w:sdt) 안까지 들어간다.

    python-docx의 iter_inner_content는 w:sdt 안의 문단·표를 건너뛰어 본문이 말없이 사라진다."""
    for child in element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from iter_blocks(content, parent)


def cell_text(cell):
    """칸의 글자를 읽는다. 칸 안의 표(중첩 표)가 있으면 그 글자도 이어 붙인다."""
    parts = [cell.text]
    for inner in cell.tables:
        for row in inner.rows:
            parts.extend(c.text for c in row.cells)
    return " ".join(p for p in parts if p.strip())


def read_table_rows(table):
    """표를 행·열 글자 목록으로 읽는다. 병합된 칸은 첫 칸에만 값을 두고 나머지는 빈 칸으로 한다.

    (행·열 목록, 빈 칸으로 바꾼 병합 칸 수, 중첩 표가 든 칸 수)를 돌려준다."""
    rows = []
    seen = set()
    merged_count = 0
    nested_count = 0
    for row in table.rows:
        cells = []
        for cell in row.cells:
            if cell._tc in seen:
                cells.append("")
                merged_count += 1
            else:
                seen.add(cell._tc)
                cells.append(cell_text(cell))
                if cell.tables:
                    nested_count += 1
        rows.append(cells)
    return rows, merged_count, nested_count


def load_docx(file_path, upload_info):
    """DOCX 본문을 순서대로 읽어 제목 스타일마다 새 섹션으로 나눈다."""
    try:
        docx_file = DocxDocument(file_path)
    except (PackageNotFoundError, BadZipFile, KeyError, XMLSyntaxError, zlib.error, AttributeError):
        # AttributeError: [Content_Types].xml만 있고 본문이 없는 zip을 python-docx가 읽다가 내는 예외 (#198)
        raise IngestError(CORRUPTED_FILE)

    doc = CommonDocument(header=dict(upload_info))
    reader = figures.FigureReader()
    chart_reader = charts.ChartReader()   # 읽은 그림·차트 수는 글자를 읽지 않은 개체 수에서 뺀다
    title = None
    lines = []
    table_no = 0
    style_names = {}
    # 제목 스타일이 하나도 없는 문서는 번호 제목(제1장·1.)이나 아웃라인 수준으로 섹션을 나눈다. 제목 스타일이 있으면 지금처럼 그것만 쓴다.
    fallback_headings = not has_heading_style(docx_file, style_names)
    fallback_count = 0

    for block in iter_blocks(docx_file.element.body, docx_file._body):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            # [병목 해결을 위해, 수업에서 다루지 않은 block._p.style 사용]
            # block.style.name은 스타일을 직접 지정하지 않은 문단마다 기본 스타일을 처음부터 다시 찾아서 느리다.
            # 그래서 스타일 번호(block._p.style)만 먼저 읽고, 처음 보는 번호일 때만 이름을 읽어 style_names에 저장해 둔다.
            style_name = style_name_of(block, style_names)
            if style_name.startswith("Heading") or (fallback_headings and looks_like_heading(text, style_name, block)):
                if not style_name.startswith("Heading"):
                    fallback_count += 1
                add_docx_section(doc, title, lines)
                title = text
                lines = []
            elif text != "":
                lines.append(text)
            lines.extend(read_pictures(docx_file, reader, block._p, lines))
            lines.extend(read_charts(docx_file, chart_reader, block._p))
        elif isinstance(block, Table):
            table_no += 1
            rows, merged_count, nested_count = read_table_rows(block)
            if merged_count > 0:
                doc.warnings.append(f"표 {table_no}: 병합 셀 {merged_count}칸은 빈 칸으로 처리함")
            if nested_count > 0:
                doc.warnings.append(f"표 {table_no}: 중첩 표가 든 {nested_count}칸은 안쪽 글자를 한 칸에 이어 붙임")
            lines.extend(table_to_text(rows))
            lines.extend(read_pictures(docx_file, reader, block._tbl, lines))
            lines.extend(read_charts(docx_file, chart_reader, block._tbl))

    add_docx_section(doc, title, lines)
    if fallback_count > 0:
        doc.warnings.append(f"제목 스타일이 없어 번호·아웃라인 제목 {fallback_count}개로 섹션을 나눔")
    unread = len(docx_file.element.body.xpath(UNREAD_OBJECTS_XPATH)) - reader.objects_read - chart_reader.appended
    if unread > 0:
        doc.warnings.append(f"그림·도형·차트·텍스트 상자 {unread}개는 글자를 읽지 않음")
    doc.warnings.extend(chart_reader.warnings())
    doc.warnings.extend(reader.warnings())
    figures.set_ocr_stats(doc, figures.ocr_status(), reader.appended, reader.unread)
    return doc