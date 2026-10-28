"""ingest/loaders/pptx_loader.py — PPTX 슬라이드 텍스트·표를 공통 섹션으로 변환.

담당: 김혜리
기준 명세: AI_문서검색_통합_기능명세서

함수:
    FUN-210 load_pptx — PPTX 슬라이드 텍스트·표를 공통 문서로 읽는다. 그림은 PaddleOCR(있으면)로, Office 차트는 차트 XML에서 읽는다."""

import zlib
from zipfile import BadZipFile

from lxml.etree import XMLSyntaxError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError
from pptx.shapes.placeholder import PlaceholderPicture

from . import charts, figures
from .base import build_location_label, table_to_text
from hyundap.common.schemas import CommonDocument, Section
from ..models import CORRUPTED_FILE, IngestError

# 글자를 읽지 않는 개체 종류 (명세 F67: 이미지·차트 속 글자 제외)
UNREAD_SHAPE_TYPES = (
    MSO_SHAPE_TYPE.PICTURE,
    MSO_SHAPE_TYPE.CHART,
    MSO_SHAPE_TYPE.MEDIA,
    MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
    MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
)
MAX_LISTED_SLIDES = 10
EMU_PER_PT = 12700
# 위쪽 끝이 이 값(EMU, 0.1인치) 이내로 다른 도형은 같은 줄로 본다. 나란히 놓은 도형의 높이가 조금 어긋나도 왼쪽부터 읽기 위함이다.
ROW_TOLERANCE_EMU = 91440


def reading_order(shapes):
    """도형을 위→아래 줄 단위로, 같은 줄 안에서는 왼→오른쪽으로 정렬한다.

    줄의 기준은 그 줄 첫 도형의 위쪽 끝이다. 도형마다 기준을 옮기면 오차가 쌓여 줄이 끝없이 이어질 수 있다."""
    rows = []   # (줄 기준 top, 도형 목록)
    for shape in sorted(shapes, key=lambda item: item.top or 0):
        top = shape.top or 0
        if rows and top - rows[-1][0] <= ROW_TOLERANCE_EMU:
            rows[-1][1].append(shape)
        else:
            rows.append((top, [shape]))
    ordered = []
    for _, row in rows:
        ordered.extend(sorted(row, key=lambda item: item.left or 0))
    return ordered


def is_picture(shape):
    """그림 도형(그림 틀 포함)인지 확인한다."""
    return shape.shape_type == MSO_SHAPE_TYPE.PICTURE or isinstance(shape, PlaceholderPicture)


def collect_shape_texts(shapes, read_picture=None, read_chart=None):
    """도형들을 위→아래, 왼→오른쪽 순서로 훑어 글자와 표 조각을 모은다. 위치 오차는 허용한다.

    read_picture(그림 도형, 지금까지 모은 글)가 있으면 그림 도형 자리에서 호출해 돌려준 글자를 모은다.
    read_chart(차트 도형)가 있으면 차트 도형 자리에서 호출해 돌려준 표 조각 목록을 모은다."""
    texts = []
    for shape in reading_order(shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts.extend(collect_shape_texts(shape.shapes, read_picture, read_chart))
        elif read_chart is not None and shape.has_chart:
            texts.extend(read_chart(shape))
        elif read_picture is not None and is_picture(shape):
            text = read_picture(shape, texts)
            if text != "":
                texts.append(text)
        elif shape.has_table:
            rows = []
            for row in shape.table.rows:
                cells = []
                for cell in row.cells:
                    cells.append(cell.text)
                rows.append(cells)
            texts.extend(table_to_text(rows))
        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text != "":
                texts.append(text)
    return texts


def count_unread_shapes(shapes):
    """글자를 읽지 않는 그림·차트·미디어·OLE 개체 수를 센다. 그룹 도형 안까지 센다."""
    count = 0
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            count += count_unread_shapes(shape.shapes)
        elif shape.shape_type in UNREAD_SHAPE_TYPES or shape.has_chart or isinstance(shape, PlaceholderPicture):
            count += 1
    return count


def load_pptx(file_path, upload_info):
    """PPTX 슬라이드를 차례로 읽어 슬라이드마다 섹션 1개로 만든다."""
    try:
        presentation = Presentation(file_path)
    except (PackageNotFoundError, BadZipFile, KeyError, XMLSyntaxError, zlib.error, AttributeError):   # 괄호 () 중 하나면 같은 처리를 한다.
        # AttributeError: [Content_Types].xml만 있고 본문이 없는 zip을 python-pptx가 읽다가 내는 예외 (#198)
        raise IngestError(CORRUPTED_FILE)

    doc = CommonDocument(header=dict(upload_info))
    reader = figures.FigureReader()
    chart_reader = charts.ChartReader()
    slide_area_pt = presentation.slide_width * presentation.slide_height / EMU_PER_PT ** 2
    read_in_slide = 0    # 지금 슬라이드에서 글자·값을 읽어 붙인 그림·차트 수(글자를 읽지 않은 개체 수에서 뺀다)
    pictures_in_slide = 0   # 그중 그림 수
    figure_slides = 0    # 그림 속 글자를 붙인 슬라이드 수

    def read_slide_picture(shape, texts):
        """그림 도형 하나에서 OCR로 읽은 글자를 돌려준다. 슬라이드에 이미 모은 글과 겹치는 줄은 뺀다."""
        nonlocal read_in_slide, pictures_in_slide
        try:
            blob = shape.image.blob
        except (KeyError, ValueError):   # 외부 연결 그림 등은 파일이 들어 있지 않다
            return ""
        width = shape.width / EMU_PER_PT if shape.width is not None else None
        height = shape.height / EMU_PER_PT if shape.height is not None else None
        text = reader.read_picture(blob, width, height, slide_area_pt, page_text=lambda: "\n".join(texts))
        if text != "":
            read_in_slide += 1
            pictures_in_slide += 1
        return text

    def read_slide_chart(shape):
        """차트 도형 하나의 제목·항목·값을 읽은 표 조각 목록을 돌려준다. 차트 파일이 없으면 건너뛴다."""
        nonlocal read_in_slide
        try:
            blob = shape.chart.part.blob
        except (KeyError, ValueError):   # 외부 연결 차트 등
            chart_reader.unreadable += 1
            return []
        pieces = chart_reader.read(blob)
        if pieces:
            read_in_slide += 1
        return pieces

    unread_slides = []   # (슬라이드 번호, 글자를 읽지 않은 개체 수)
    for number, slide in enumerate(presentation.slides, start=1):
        read_in_slide = pictures_in_slide = 0
        texts = collect_shape_texts(slide.shapes, read_slide_picture, read_slide_chart)
        figure_slides += 1 if pictures_in_slide > 0 else 0
        unread = count_unread_shapes(slide.shapes) - read_in_slide
        if unread > 0:
            unread_slides.append((number, unread))
        if len(texts) == 0:
            # 발표자 노트는 읽지 않는다(명세). 노트에만 글자가 있으면 그 사실을 경고에 구분해 남긴다.
            notes_frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
            has_notes_text = notes_frame is not None and notes_frame.text.strip() != ""
            reason = "본문 글자가 없어 건너뜀(발표자 노트에만 글자가 있음, 노트는 제외)" if has_notes_text else "글자가 없어 건너뜀"
            doc.warnings.append(f"슬라이드 {number}: {reason}")
            continue
        doc.sections.append(
            Section(
                section_no=number,
                location_label=build_location_label("slide", number),
                text="\n\n".join(texts),
            )
        )
    if unread_slides:
        # 슬라이드가 많은 문서에서 경고가 늘어나지 않도록 문서당 1개로 모은다.
        listed = ", ".join(str(number) for number, _ in unread_slides[:MAX_LISTED_SLIDES])
        more = " 외" if len(unread_slides) > MAX_LISTED_SLIDES else ""
        total = sum(count for _, count in unread_slides)
        doc.warnings.append(f"이미지·차트 {total}개는 글자를 읽지 않음 (슬라이드 {listed}{more})")
    doc.warnings.extend(chart_reader.warnings())
    doc.warnings.extend(reader.warnings(f"슬라이드 {figure_slides}장, "))
    return doc