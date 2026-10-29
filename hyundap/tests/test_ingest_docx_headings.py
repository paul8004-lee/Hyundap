"""제목 스타일(Heading)이 없는 DOCX의 번호·아웃라인 제목 섹션 나누기 테스트."""
import unittest

from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from hyundap.ingest.loaders.docx_loader import HEADING_MAX_CHARS, load_docx, looks_like_heading
from hyundap.tests.test_ingest import INFO, IngestTestCase


def labels(doc):
    return [section.location_label for section in doc.sections]


class FallbackHeadingTests(IngestTestCase):
    def load(self, document):
        document.save(self.path('h.docx'))
        return load_docx(self.path('h.docx'), dict(INFO))

    def test_numbered_titles_split_sections_without_heading_styles(self):
        document = DocxDocument()
        document.add_paragraph('보고서 제목')
        for number, title in enumerate(['사업 개요', '사업 내용', '사업 평가'], start=1):
            document.add_paragraph(f'{number}. {title}')
            document.add_paragraph(f'{title} 본문입니다.')
        doc = self.load(document)
        self.assertEqual(labels(doc), ['섹션 1', '섹션 2 (1. 사업 개요)', '섹션 3 (2. 사업 내용)', '섹션 4 (3. 사업 평가)'])
        self.assertEqual(doc.sections[1].text, '1. 사업 개요\n\n사업 개요 본문입니다.')
        self.assertIn('제목 스타일이 없어 번호·아웃라인 제목 3개로 섹션을 나눔', doc.warnings)

    def test_korean_chapter_and_roman_titles(self):
        document = DocxDocument()
        document.add_paragraph('제1장 총칙')
        document.add_paragraph('내용')
        document.add_paragraph('Ⅱ. 현황')
        document.add_paragraph('내용')
        self.assertEqual(labels(self.load(document)), ['섹션 1 (제1장 총칙)', '섹션 2 (Ⅱ. 현황)'])

    def test_small_numbers_stay_in_body(self):
        """'1)'·'가.' 같은 작은 번호는 제목으로 보지 않아 작은 섹션이 늘지 않는다."""
        document = DocxDocument()
        document.add_paragraph('1. 사업 내용')
        document.add_paragraph('1) 진행 경과')
        document.add_paragraph('2) 세부 내용')
        doc = self.load(document)
        self.assertEqual(labels(doc), ['섹션 1 (1. 사업 내용)'])
        self.assertIn('1) 진행 경과', doc.sections[0].text)

    def test_long_sentences_and_dates_are_not_titles(self):
        document = DocxDocument()
        document.add_paragraph('2022.08.23')
        document.add_paragraph('3. 이 문단은 번호로 시작하지만 문장이므로 제목이 아니며 길이도 제목보다 훨씬 길게 이어집니다.')
        document.add_paragraph('4. 행사를 진행했다.')
        document.add_paragraph('10.5 억원이 집행됨')
        doc = self.load(document)
        self.assertEqual(labels(doc), ['섹션 1'])
        self.assertNotIn('제목 스타일이 없어', ' '.join(doc.warnings))

    def test_title_ending_with_yo_is_still_a_title(self):
        """'개요'처럼 '요'로 끝나는 제목은 문장으로 오인하지 않는다."""
        document = DocxDocument()
        document.add_paragraph('1. 사업 개요')
        document.add_paragraph('내용')
        self.assertEqual(labels(self.load(document)), ['섹션 1 (1. 사업 개요)'])

    def test_list_paragraphs_are_not_titles(self):
        document = DocxDocument()
        document.add_paragraph('1. 항목 하나', style='List Paragraph')
        document.add_paragraph('본문')
        self.assertEqual(labels(self.load(document)), ['섹션 1'])

    def test_table_cell_numbers_are_not_titles(self):
        document = DocxDocument()
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = '1. 표 안의 번호'
        document.add_paragraph('본문')
        self.assertEqual(labels(self.load(document)), ['섹션 1'])

    def test_outline_level_paragraph_is_a_title(self):
        document = DocxDocument()
        document.add_paragraph('앞 내용')
        paragraph = document.add_paragraph('개요')
        outline = OxmlElement('w:outlineLvl')
        outline.set(qn('w:val'), '0')
        paragraph._p.get_or_add_pPr().append(outline)
        document.add_paragraph('뒤 내용')
        self.assertEqual(labels(self.load(document)), ['섹션 1', '섹션 2 (개요)'])

    def test_heading_styles_keep_current_behavior(self):
        """제목 스타일이 하나라도 있으면 번호만 붙은 일반 문단은 제목으로 보지 않는다(기존 동작)."""
        document = DocxDocument()
        document.add_heading('개요', level=1)
        document.add_paragraph('1. 번호만 붙은 일반 문단')
        document.add_heading('결론', level=1)
        doc = self.load(document)
        self.assertEqual(labels(doc), ['섹션 1 (개요)', '섹션 2 (결론)'])
        self.assertNotIn('제목 스타일이 없어', ' '.join(doc.warnings))

    def test_rule_boundary(self):
        class Style:
            _p = type('P', (), {'pPr': None})()
        self.assertTrue(looks_like_heading('1. ' + '가' * (HEADING_MAX_CHARS - 3), 'Normal', Style))
        self.assertFalse(looks_like_heading('1. ' + '가' * (HEADING_MAX_CHARS - 2), 'Normal', Style))
        self.assertFalse(looks_like_heading('', 'Normal', Style))


if __name__ == '__main__':
    unittest.main()
