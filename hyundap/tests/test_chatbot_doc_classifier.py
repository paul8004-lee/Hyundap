"""chatbot/doc_classifier.py 테스트. OpenAI·DB·Chroma를 부르지 않고 대역으로 실행한다.

실행(저장소 맨 위 폴더에서):  python -m unittest hyundap.tests.test_chatbot_doc_classifier -v
"""
import unittest
from unittest.mock import patch

from openai.types.decision import Decision

from hyundap.chatbot import doc_classifier
from hyundap.chatbot.constants import CLASSIFY_MODEL, CLASSIFY_SAMPLE_CHARS
from hyundap.chatbot.models import ChatbotError, DocClassification
from hyundap.common.schemas import CommonDocument, Section


# ingest 정규화가 끝난 문서 (업로드 경로 입력)
def common_doc(doc_type='manual', sections=None):
    if sections is None:
        sections = [Section(1, 'p.1', 'A사 사내 규정 제1조'), Section(2, 'p.2', '제2조 책임')]
    return CommonDocument(header={'document_id': 'd9', 'version_id': 'v9', 'company_id': 'co', 'company_name': 'A사',
                                  'doc_type': doc_type, 'file_name': '규정.docx'}, sections=sections)


# Decisions API 응답을 SDK의 실제 응답 형식(Decision)으로 만든다
def decision(*answers):
    return Decision.model_validate({
        'answers': list(answers), 'model': 'gpt-6-luna',
        'usage': {'input_tokens': 400, 'output_tokens': 0, 'total_tokens': 400,
                  'input_tokens_details': {'cached_tokens': 0, 'cache_write_tokens': 0},
                  'output_tokens_details': {'reasoning_tokens': 0}},
    })


def choice(value, confidence=0.95):
    return {'type': 'choice', 'name': 'doc_type', 'choice': value, 'confidence': confidence,
            'probabilities': [{'value': value, 'probability': confidence}]}


def predicate(name, probability):
    return {'type': 'predicate', 'name': name, 'probability': probability}


GOOD = decision(choice('policy', 0.93), predicate('company_named', 0.98), predicate('company_match', 0.96))


class FakeClient:
    """client.decisions.create 대역. reply가 예외면 던지고 아니면 그대로 돌려준다."""

    def __init__(self, reply):
        self.reply, self.calls = reply, []
        self.decisions = self

    def create(self, **kwargs):
        self.calls.append(kwargs)
        if isinstance(self.reply, Exception):
            raise self.reply
        return self.reply


def run_common(reply, doc=None):
    """업로드 경로(classify_common_document)를 대역 클라이언트로 실행하고 (결과, 클라이언트)를 돌려준다."""
    client = FakeClient(reply)
    with patch.object(doc_classifier, 'get_classify_client', return_value=client):
        result = doc_classifier.classify_common_document(common_doc() if doc is None else doc)
    return result, client


# 입력·질문 만들기와 응답 읽기 (호출 없음)
class Pieces(unittest.TestCase):
    def test_input_uses_ordered_sample_within_limit(self):
        text = doc_classifier.build_classify_input('점검절차.docx', ['첫째' * 10, '  ', '둘째' * 10, '가' * CLASSIFY_SAMPLE_CHARS])
        self.assertTrue(text.startswith('파일명: 점검절차.docx'))
        self.assertLess(text.index('첫째'), text.index('둘째'))
        self.assertEqual(len(text.split('본문:\n')[1].replace('\n', '')), CLASSIFY_SAMPLE_CHARS)
        # 업로드 경로는 정규화된 섹션 순서대로 본문을 넘긴다
        _, client = run_common(GOOD, common_doc(sections=[Section(1, 'p.1', '첫째'), Section(2, 'p.2', '둘째')]))
        sent = client.calls[0]['input']
        self.assertLess(sent.index('첫째'), sent.index('둘째'))

    def test_questions(self):
        questions = doc_classifier.build_classify_questions('A사')
        self.assertEqual([q['name'] for q in questions], ['doc_type', 'company_named', 'company_match'])
        self.assertEqual([c['value'] for c in questions[0]['choices']], ['manual', 'policy', 'guide', 'other'])
        self.assertIn("'A사'", questions[2]['instructions'])
        self.assertEqual([q['name'] for q in doc_classifier.build_classify_questions('  ')], ['doc_type'])

    def test_read_answers_company_rules(self):
        cases = [((0.98, 0.96), 'yes'), ((0.98, 0.05), 'no'), ((0.98, 0.5), ''), ((0.1, 0.99), '')]
        for (named, match), expected in cases:
            with self.subTest(named=named, match=match):
                answer = doc_classifier.read_classify_answers(decision(
                    choice('guide', 0.8), predicate('company_named', named), predicate('company_match', match)))
                self.assertEqual(answer, ('guide', 0.8, expected))

    def test_read_answers_unknown_type_and_refusal(self):
        self.assertEqual(doc_classifier.read_classify_answers(decision(choice('contract', 0.9))), ('', 0.0, ''))
        with self.assertRaises(ValueError):
            doc_classifier.read_classify_answers(decision({'type': 'refusal', 'name': 'doc_type'}))
        # 회사 질문만 거부되면 유형은 남기고 회사 일치는 비운다
        answer = doc_classifier.read_classify_answers(decision(
            choice('manual'), {'type': 'refusal', 'name': 'company_match'}, predicate('company_named', 0.9)))
        self.assertEqual(answer, ('manual', 0.95, ''))


# 업로드 경로 분류 흐름: 정규화된 문서 → API 1회 (접근 확인은 업로드 쪽이 먼저 한다)
class Flow(unittest.TestCase):
    def test_success_compares_with_header_type(self):
        result, client = run_common(GOOD)
        self.assertEqual(result, DocClassification(
            document_id='d9', version_id='v9', admin_doc_type='manual', auto_doc_type='policy', confidence=0.93,
            matches_admin_type=False, company_match='yes', status='success', detail=''))
        call = client.calls[0]
        self.assertEqual(call['model'], CLASSIFY_MODEL)
        self.assertIn('파일명: 규정.docx', call['input'])
        self.assertIn('A사 사내 규정 제1조', call['input'])
        self.assertEqual(len(call['questions']), 3)
        self.assertIn("'A사'", call['questions'][2]['instructions'])

    def test_same_type_marks_match(self):
        result, _ = run_common(decision(choice('manual', 0.9)))
        self.assertIs(result.matches_admin_type, True)

    def test_no_access_check_and_no_call_without_text(self):
        with patch('hyundap.auth.permissions.verify_documents_access') as access:
            run_common(GOOD)
        access.assert_not_called()
        # 본문이 없으면 API를 부르지 않는다
        for doc in (common_doc(sections=[]), object()):
            with self.subTest(doc=type(doc).__name__):
                with patch.object(doc_classifier, 'get_classify_client') as get_client:
                    result = doc_classifier.classify_common_document(doc)
                self.assertEqual((result.status, result.detail), ('failed', 'EMPTY_DOCUMENT'))
                get_client.assert_not_called()

    def test_api_failure_returns_failed_without_secret(self):
        for reply in (RuntimeError('boom sk-test-key'), TimeoutError('slow'),
                      decision({'type': 'refusal', 'name': 'doc_type'})):
            with self.subTest(reply=type(reply).__name__):
                result, _ = run_common(reply)
                self.assertEqual((result.status, result.auto_doc_type, result.matches_admin_type), ('failed', '', None))
                self.assertNotIn('sk-test', result.detail)
        # 클라이언트를 만들다 실패해도 올리지 않는다
        with patch.object(doc_classifier, 'get_classify_client', side_effect=RuntimeError('down')):
            self.assertEqual(doc_classifier.classify_common_document(common_doc()).status, 'failed')

    def test_missing_key_is_failed_not_raised(self):
        with patch('hyundap.common.config.get_settings') as settings:
            settings.return_value.openai_api_key = ''
            with self.assertRaises(ChatbotError) as caught:
                doc_classifier.get_classify_client()
        self.assertEqual(caught.exception.code, 'CONFIG_MISSING')
        # 업로드 경로는 키 없음도 예외 대신 failed로 돌려준다
        with patch.object(doc_classifier, 'get_classify_client', side_effect=ChatbotError('CONFIG_MISSING')):
            result = doc_classifier.classify_common_document(common_doc())
        self.assertEqual((result.status, result.detail), ('failed', 'CONFIG_MISSING'))

    def test_real_client_settings(self):
        doc_classifier._CLIENTS.clear()
        self.addCleanup(doc_classifier._CLIENTS.clear)
        with patch('hyundap.common.config.get_settings') as settings:
            settings.return_value.openai_api_key = 'sk-test-key'
            client = doc_classifier.get_classify_client()
            self.assertIs(client, doc_classifier.get_classify_client())
        self.assertEqual((client.timeout, client.max_retries), (10, 0))
        self.assertTrue(hasattr(client.decisions, 'create'))


# 자동 분류 업로드의 최종 유형 규칙 (#261 확정안 1-2): (doc_type, doc_type_fallback, doc_type_fallback_reason)
class Decide(unittest.TestCase):
    def classified(self, auto_doc_type='policy', confidence=0.93, status='success'):
        return DocClassification('d1', 'v1', 'other', auto_doc_type, confidence, None, 'yes', status)

    def test_confident_type_is_used_for_result_and_report(self):
        # 분류기 결과(DocClassification)와 수집 보고서 dict를 모두 받는다
        report = {'auto_doc_type': 'policy', 'doc_type_confidence': 0.93, 'company_match': 'yes',
                  'classify_status': 'success'}
        for value in (self.classified(), report):
            with self.subTest(kind=type(value).__name__):
                self.assertEqual(doc_classifier.decide_doc_type(value, 'other'), ('policy', False, None))

    def test_threshold_boundary(self):
        decide = doc_classifier.decide_doc_type
        self.assertEqual(decide(self.classified(confidence=0.8), 'other'), ('policy', False, None))
        self.assertEqual(decide(self.classified(confidence=0.799), 'other'), ('other', True, 'LOW_CONFIDENCE'))
        # 기준은 상수 하나로 바뀐다
        with patch.object(doc_classifier, 'CLASSIFY_APPLY_MIN_CONFIDENCE', 0.95):
            self.assertEqual(decide(self.classified(confidence=0.93), 'other'), ('other', True, 'LOW_CONFIDENCE'))

    def test_ai_other_is_not_fallback(self):
        # AI가 기준 이상으로 '기타'라고 판단하면 대체 저장이 아니다
        self.assertEqual(doc_classifier.decide_doc_type(self.classified('other', 0.9), 'other'), ('other', False, None))

    def test_failed_or_unusable_result_is_classify_failed(self):
        cases = [
            self.classified('', 0.0, 'failed'),           # API 실패·거부 응답
            self.classified('', 0.0),                     # 목록 밖 유형 (read_classify_answers가 ''로 비움)
            self.classified('policy', 0.93, 'failed'),    # 상태가 실패면 값이 있어도 쓰지 않는다
            {'classify_status': 'failed'},                # 수집이 분류기 예외를 받은 경우
            {}, None, 'policy',
            {'classify_status': 'success', 'auto_doc_type': 'contract', 'doc_type_confidence': 0.9},
            {'classify_status': 'success', 'auto_doc_type': 'policy'},
            {'classify_status': 'success', 'auto_doc_type': 'policy', 'doc_type_confidence': True},
            {'classify_status': 'success', 'auto_doc_type': 'policy', 'doc_type_confidence': '0.9'},
            {'classify_status': 'success', 'auto_doc_type': 'policy', 'doc_type_confidence': float('nan')},
            {'classify_status': 'success', 'auto_doc_type': 'policy', 'doc_type_confidence': 1.5},
        ]
        for value in cases:
            with self.subTest(value=value):
                self.assertEqual(doc_classifier.decide_doc_type(value, 'other'), ('other', True, 'CLASSIFY_FAILED'))

    def test_new_version_keeps_current_type(self):
        # 새 버전은 현재 유형이 기존 유형이다. 대체 저장이면 기존 유형을 유지하고, 확정이면 AI 유형으로 바뀐다
        decide = doc_classifier.decide_doc_type
        self.assertEqual(decide(self.classified('manual', 0.6), 'policy'), ('policy', True, 'LOW_CONFIDENCE'))
        self.assertEqual(decide(self.classified('', 0.0, 'failed'), 'policy'), ('policy', True, 'CLASSIFY_FAILED'))
        self.assertEqual(decide(self.classified('manual', 0.9), 'policy'), ('manual', False, None))

    def test_bad_current_type_uses_default(self):
        for current in (None, '', 'contract', ['policy']):
            with self.subTest(current=current):
                self.assertEqual(doc_classifier.decide_doc_type(self.classified(confidence=0.5), current),
                                 ('other', True, 'LOW_CONFIDENCE'))


# 업로드 연결: ingest 정규화 직후 분류 → 수집 보고서 → 작업 기록(jobs_repo). API는 부르지 않는다
class UploadPath(unittest.TestCase):
    def make_docx(self, text):
        import tempfile
        from pathlib import Path
        from docx import Document
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        path = Path(folder.name) / '규정.docx'
        document = Document()
        document.add_paragraph(text)
        document.save(path)
        return str(path)

    def ingest(self, **kwargs):
        from hyundap.ingest.pipeline import ingest_document
        info = {'document_id': 'd1', 'version_id': 'v1', 'company_id': 'co', 'company_name': 'A사',
                'doc_type': 'manual', 'file_name': '규정.docx'}
        return ingest_document(self.make_docx('A사 사내 규정 제1조 이 규정은 연구윤리 기준을 정한다.'), info, **kwargs)

    def test_ingest_classifies_only_when_asked(self):
        good = DocClassification('d1', 'v1', 'manual', 'policy', 0.93, False, 'yes', 'success')
        with patch.object(doc_classifier, 'classify_common_document', return_value=good) as classify:
            plain = self.ingest()
            classified = self.ingest(classify=True)
        self.assertEqual(classify.call_count, 1)            # 기본(classify=False)은 부르지 않는다
        self.assertNotIn('auto_doc_type', plain.report)
        doc = classify.call_args.args[0]
        self.assertEqual((doc.header['doc_type'], len(doc.sections) > 0), ('manual', True))   # 정규화된 문서를 받는다
        self.assertEqual({key: classified.report[key] for key in
                          ('auto_doc_type', 'doc_type_confidence', 'company_match', 'classify_status')},
                         {'auto_doc_type': 'policy', 'doc_type_confidence': 0.93, 'company_match': 'yes',
                          'classify_status': 'success'})
        # 청크와 관리자 doc_type은 그대로다
        self.assertEqual([c.text for c in classified.chunks], [c.text for c in plain.chunks])
        self.assertEqual({c.doc_type for c in classified.chunks}, {'manual'})

    def test_ingest_continues_when_classifier_breaks(self):
        with patch.object(doc_classifier, 'classify_common_document', side_effect=RuntimeError('boom sk-test-key')), \
                self.assertLogs('hyundap.ingest.pipeline', 'WARNING') as logs:
            result = self.ingest(classify=True)
        self.assertEqual((result.status, result.report['classify_status']), ('success', 'failed'))
        self.assertNotIn('auto_doc_type', result.report)
        # 실패 원인은 예외 이름만 로그에 남고, 오류 문구(키가 섞일 수 있음)는 남지 않는다
        self.assertIn('RuntimeError', logs.output[0])
        self.assertNotIn('sk-test-key', ''.join(logs.output))

    def test_ingest_logs_cause_when_classification_fails(self):
        # 키 누락·모델 권한 문제 등으로 분류기가 failed를 돌려주면 원인 이름을 로그로 남긴다
        failed = DocClassification('d1', 'v1', 'manual', status='failed', detail='PermissionDeniedError')
        with patch.object(doc_classifier, 'classify_common_document', return_value=failed), \
                self.assertLogs('hyundap.ingest.pipeline', 'WARNING') as logs:
            result = self.ingest(classify=True)
        self.assertEqual((result.status, result.report['classify_status']), ('success', 'failed'))
        self.assertIn('PermissionDeniedError', logs.output[0])
        # 성공하면 경고를 남기지 않는다
        good = DocClassification('d1', 'v1', 'manual', 'policy', 0.93, False, 'yes', 'success')
        with patch.object(doc_classifier, 'classify_common_document', return_value=good), \
                self.assertNoLogs('hyundap.ingest.pipeline', 'WARNING'):
            self.ingest(classify=True)

    def test_upload_turns_on_classify_only_for_supporting_ingest(self):
        import os
        from hyundap.common import config
        from hyundap.db import upload_service, version_service
        seen = []

        def capture(function, *args, **kwargs):
            seen.append(kwargs)
            return {'status': 'failed', 'error_code': 'EMPTY_DOCUMENT'}

        def new_ingest(path, info, find_by_hash=None, file_hash=None, classify=False):
            pass

        def old_ingest(path, info, find_by_hash=None):
            pass

        ids = {'document_id': 'd1', 'version_id': 'v1', 'job_id': 'j1'}
        try:
            with patch.object(config, 'load_dotenv'), patch.object(version_service, 'call_external', side_effect=capture):
                for value, function in ((None, new_ingest), ('0', new_ingest), ('1', new_ingest), ('1', old_ingest)):
                    env = {} if value is None else {'DOC_CLASSIFY': value}
                    with patch.dict(os.environ, env, clear=True):
                        config.get_settings.cache_clear()
                        upload_service.ingest_and_save_chunks('a.docx', {'company_id': 'co'}, ids, function)
            self.assertNotIn('classify', seen[0])
            self.assertNotIn('classify', seen[1])
            self.assertIs(seen[2].get('classify'), True)
            self.assertNotIn('classify', seen[3])  # 옛 수집기에는 넘기지 않는다.
        finally:
            config.get_settings.cache_clear()

    def test_job_report_keeps_ingest_classification(self):
        from hyundap.db import jobs_repo
        from hyundap.db.errors import DbError
        base = dict(stage='ingest', status='success', parser='python-docx', chunk_count=1)
        clean = jobs_repo.normalize_report(dict(base, auto_doc_type='policy', doc_type_confidence=0.93,
                                                company_match='yes', classify_status='success'))
        self.assertEqual({key: clean[key] for key in jobs_repo.CLASSIFY_FIELDS},
                         {'auto_doc_type': 'policy', 'doc_type_confidence': 0.93, 'company_match': 'yes',
                          'classify_status': 'success'})
        self.assertEqual(jobs_repo.normalize_report(dict(base, classify_status='failed'))['classify_status'], 'failed')
        # 범위 밖 값은 수집 기록을 거부하지 않고 분류 값만 빼고 failed로 남긴다 (업로드는 성공)
        good = dict(auto_doc_type='policy', doc_type_confidence=0.93, company_match='yes', classify_status='success')
        bad = [dict(auto_doc_type='contract'), dict(auto_doc_type=None), dict(doc_type_confidence=1.5),
               dict(doc_type_confidence=True), dict(doc_type_confidence=float('nan')), dict(doc_type_confidence='0.9'),
               dict(company_match='maybe'), dict(classify_status='skipped'), dict(classify_status=None)]
        for changes in bad:
            with self.subTest(changes=changes):
                clean = jobs_repo.normalize_report(dict(base, **dict(good, **changes)))
                self.assertEqual({key: clean[key] for key in jobs_repo.CLASSIFY_FIELDS if key in clean},
                                 {'auto_doc_type': '', 'doc_type_confidence': 0.0,
                                  'company_match': '', 'classify_status': 'failed'})
                self.assertEqual((clean['stage'], clean['status'], clean['chunk_count']), ('ingest', 'success', 1))
        # 상태 없이 값만 오거나 상태만 이상해도 failed
        self.assertEqual(jobs_repo.normalize_report(dict(base, auto_doc_type='policy'))['classify_status'], 'failed')
        # 분류 값이 없으면 아무것도 붙이지 않는다
        self.assertNotIn('classify_status', jobs_repo.normalize_report(base))
        with self.assertRaises(DbError):   # 다른 단계에는 섞일 수 없다
            jobs_repo.normalize_report(dict(stage='embedding', status='success', auto_doc_type='policy'))


if __name__ == '__main__':
    unittest.main()
