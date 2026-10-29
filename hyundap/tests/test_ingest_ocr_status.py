"""수집 보고서의 OCR 상태(ocr_status)·읽은 수(ocr_read_count)·읽지 못한 수(ocr_unread_count)와 OCR 꺼짐 경고 테스트."""
import io
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document as DocxDocument
from docx.shared import Inches
from langchain_core.documents import Document
from pptx import Presentation
from pptx.util import Inches as PptxInches

from hyundap.ingest.loaders.docx_loader import load_docx
from hyundap.ingest.loaders.pdf_loader import load_pdf
from hyundap.ingest.loaders.pptx_loader import load_pptx
from hyundap.ingest.normalizer import normalize_document
from hyundap.tests.test_ingest import INFO, PDFIUM, IngestTestCase, ocr_settings, pdf_with_images, png_bytes, minimal_pdf

UNAVAILABLE = 'OCR이 설치되어 있지 않아 이미지·스캔 쪽 {}개의 글자를 읽지 않음'
DISABLED = 'OCR이 설정으로 꺼져 있어 이미지·스캔 쪽 {}개의 글자를 읽지 않음'


def stats(doc):
    return doc.ocr['status'], doc.ocr['read'], doc.ocr['unread']


class PdfOcrStatusTests(IngestTestCase):
    def figure_pdf(self):
        Path(self.path('fig.pdf')).write_bytes(pdf_with_images('Hello', [(200, 150)]))
        return self.path('fig.pdf')

    @PDFIUM
    def test_unavailable_counts_unread_figure_and_warns_with_reason(self):
        """OCR가 설치되어 있지 않으면 글자 있는 쪽의 그림을 읽지 못한 것으로 세고 이유가 담긴 경고를 남긴다."""
        with patch('hyundap.ingest.ocr.is_available', return_value=False):
            doc = load_pdf(self.figure_pdf(), dict(INFO))
        self.assertEqual(stats(doc), ('unavailable', 0, 1))
        self.assertEqual(doc.warnings, [UNAVAILABLE.format(1)])

    @PDFIUM
    def test_enabled_counts_read_figure_without_off_warning(self):
        with patch('hyundap.ingest.ocr.is_available', return_value=True), \
                patch('hyundap.ingest.ocr.ocr_image', return_value='그림 글자'):
            doc = load_pdf(self.figure_pdf(), dict(INFO))
        self.assertEqual(stats(doc), ('enabled', 1, 0))
        self.assertFalse([w for w in doc.warnings if '읽지 않음' in w])

    @PDFIUM
    def test_disabled_by_settings_is_reported_as_disabled(self):
        with patch('hyundap.ingest.ocr.is_available', return_value=True), \
                ocr_settings(OCR_MAX_PAGES=0, OCR_MAX_FIGURES=0):
            doc = load_pdf(self.figure_pdf(), dict(INFO))
        self.assertEqual(stats(doc), ('disabled', 0, 1))
        self.assertEqual(doc.warnings, [DISABLED.format(1)])

    @PDFIUM
    def test_only_figure_limit_zero_stays_enabled_for_pdf_pages(self):
        """PDF는 쪽 OCR 상한이 0보다 크면 그림 상한이 0이어도 켜짐으로 본다. 못 읽은 그림은 개수에 남는다."""
        with patch('hyundap.ingest.ocr.is_available', return_value=True), \
                ocr_settings(OCR_MAX_PAGES=5, OCR_MAX_FIGURES=0):
            doc = load_pdf(self.figure_pdf(), dict(INFO))
        self.assertEqual(stats(doc), ('enabled', 0, 1))
        self.assertFalse([w for w in doc.warnings if 'OCR이' in w])

    def test_blank_pages_count_as_unread_when_ocr_is_off(self):
        """글자 없는 쪽은 OCR을 쓸 수 없으면 읽지 못한 쪽으로 센다."""
        pages = [Document(page_content='본문', metadata={'page': 0}), Document(page_content='', metadata={'page': 1})]
        with patch('hyundap.ingest.ocr.is_available', return_value=False), \
                patch('hyundap.ingest.loaders.pdf_loader.PyPDFLoader') as loader:
            loader.return_value.load.return_value = pages
            doc = load_pdf(self.path('x.pdf'), dict(INFO))
        self.assertEqual(stats(doc), ('unavailable', 0, 1))
        self.assertEqual(doc.warnings, ['p.2: 글자가 없어 건너뜀', UNAVAILABLE.format(1)])

    @PDFIUM
    def test_pages_over_limit_are_unread_but_status_stays_enabled(self):
        """쪽 수 상한을 넘은 쪽은 읽지 못한 쪽으로 세지만 OCR 상태는 켜짐이고 '꺼져 있어' 경고는 없다."""
        pages = [Document(page_content='', metadata={'page': n}) for n in range(3)]
        Path(self.path('scan.pdf')).write_bytes(minimal_pdf(''))
        with patch('hyundap.ingest.ocr.is_available', return_value=True), \
                patch('hyundap.ingest.ocr.ocr_image', return_value='본문'), \
                patch('hyundap.ingest.loaders.pdf_loader.render_page'), \
                ocr_settings(OCR_MAX_PAGES=1), \
                patch('hyundap.ingest.loaders.pdf_loader.PyPDFLoader') as loader:
            loader.return_value.load.return_value = pages
            doc = load_pdf(self.path('scan.pdf'), dict(INFO))
        self.assertEqual(stats(doc), ('enabled', 1, 2))
        self.assertFalse([w for w in doc.warnings if 'OCR이' in w])

    @PDFIUM
    def test_text_pdf_without_figures_reports_zero_counts(self):
        Path(self.path('a.pdf')).write_bytes(minimal_pdf('Hello'))
        with patch('hyundap.ingest.ocr.is_available', return_value=False):
            doc = load_pdf(self.path('a.pdf'), dict(INFO))
        self.assertEqual(stats(doc), ('unavailable', 0, 0))
        self.assertEqual(doc.warnings, [])

    @PDFIUM
    def test_counting_failure_does_not_fail_the_upload(self):
        with patch('hyundap.ingest.ocr.is_available', return_value=False), \
                patch('hyundap.ingest.loaders.pdf_loader.count_figures', side_effect=RuntimeError('boom')):
            doc = load_pdf(self.figure_pdf(), dict(INFO))
        self.assertEqual(stats(doc), ('unavailable', 0, 0))


class OfficeOcrStatusTests(IngestTestCase):
    def docx_with_picture(self, width):
        document = DocxDocument()
        document.add_paragraph('본문')
        document.add_picture(io.BytesIO(png_bytes()), width=width)
        document.save(self.path('p.docx'))
        return self.path('p.docx')

    def test_docx_picture_unread_when_ocr_unavailable(self):
        with patch('hyundap.ingest.ocr.is_available', return_value=False):
            doc = load_docx(self.docx_with_picture(Inches(3)), dict(INFO))
        self.assertEqual(stats(doc), ('unavailable', 0, 1))
        self.assertIn(UNAVAILABLE.format(1), doc.warnings)

    def test_docx_small_picture_is_not_counted(self):
        """로고처럼 읽을 크기가 아닌 그림은 읽지 못한 그림으로 세지 않는다."""
        with patch('hyundap.ingest.ocr.is_available', return_value=False):
            doc = load_docx(self.docx_with_picture(Inches(0.3)), dict(INFO))
        self.assertEqual(stats(doc), ('unavailable', 0, 0))
        self.assertFalse([w for w in doc.warnings if 'OCR이' in w])

    def test_docx_enabled_counts_read_picture(self):
        with patch('hyundap.ingest.ocr.is_available', return_value=True), \
                patch('hyundap.ingest.ocr.ocr_image', return_value='그림 글자'):
            doc = load_docx(self.docx_with_picture(Inches(3)), dict(INFO))
        self.assertEqual(stats(doc), ('enabled', 1, 0))

    def test_docx_disabled_by_settings(self):
        with patch('hyundap.ingest.ocr.is_available', return_value=True), ocr_settings(OCR_MAX_FIGURES=0):
            doc = load_docx(self.docx_with_picture(Inches(3)), dict(INFO))
        self.assertEqual(stats(doc), ('disabled', 0, 1))
        self.assertIn(DISABLED.format(1), doc.warnings)

    def test_pptx_picture_unread_when_ocr_unavailable(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = '제목'
        slide.shapes.add_picture(io.BytesIO(png_bytes()), PptxInches(1), PptxInches(2), width=PptxInches(3))
        presentation.save(self.path('p.pptx'))
        with patch('hyundap.ingest.ocr.is_available', return_value=False):
            doc = load_pptx(self.path('p.pptx'), dict(INFO))
        self.assertEqual(stats(doc), ('unavailable', 0, 1))
        self.assertIn(UNAVAILABLE.format(1), doc.warnings)

    def test_unreadable_format_picture_counts_as_unread(self):
        """읽을 수 없는 형식의 그림은 OCR이 켜져 있어도 읽지 못한 그림으로 센다."""
        with patch('hyundap.ingest.ocr.is_available', return_value=True), \
                patch('hyundap.ingest.loaders.figures.picture_to_array', side_effect=__import__(
                    'hyundap.ingest.loaders.figures', fromlist=['UnreadablePicture']).UnreadablePicture('emf')):
            doc = load_docx(self.docx_with_picture(Inches(3)), dict(INFO))
        self.assertEqual(stats(doc), ('enabled', 0, 1))


class ReportKeyTests(IngestTestCase):
    def test_report_has_ocr_keys_and_normalizer_keeps_them(self):
        document = DocxDocument()
        document.add_paragraph('본문')
        document.save(self.path('a.docx'))
        with patch('hyundap.ingest.ocr.is_available', return_value=False):
            result = self.ingest('a.docx')
        self.assertEqual((result.report['ocr_status'], result.report['ocr_read_count'], result.report['ocr_unread_count']),
                         ('unavailable', 0, 0))
        doc = load_docx(self.path('a.docx'), dict(INFO))
        self.assertEqual(normalize_document(doc).ocr, doc.ocr)


if __name__ == '__main__':
    unittest.main()
